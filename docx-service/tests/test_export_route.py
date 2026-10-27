import io

from docx import Document
from fastapi.testclient import TestClient

import main

client = TestClient(main.app)
_auth_headers = {"X-Internal-Api-Key": main._internal_api_key}


def _profile_payload() -> dict:
    return {
        "id": "1",
        "name": "Test Profile",
        "header": {
            "fullName": "Jane Doe",
            "careerTitle": "AI Engineer",
            "email": "jane@example.com",
            "phone": "555-0100",
            "location": "Remote",
            "links": [],
            "primaryColor": "#2563eb",
            "secondaryColor": "#7c3aed",
        },
        "sections": [
            {
                "id": "s1",
                "type": "text",
                "title": "Summary",
                "body": "**Bold** summary.",
            }
        ],
    }


def test_export_returns_a_valid_docx_with_the_profile_content() -> None:
    response = client.post("/export", json=_profile_payload(), headers=_auth_headers)

    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    assert "Jane-Doe_Test-Profile_CV.docx" in response.headers["content-disposition"]

    document = Document(io.BytesIO(response.content))
    header_texts = [p.text for p in document.tables[0].rows[0].cells[-1].paragraphs]
    assert "Jane Doe" in header_texts
    assert any(p.text == "Summary" for p in document.paragraphs)


def test_export_rejects_a_request_without_the_internal_api_key() -> None:
    response = client.post("/export", json=_profile_payload())

    assert response.status_code == 401

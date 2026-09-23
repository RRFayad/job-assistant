import base64
import io

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from schemas.profile import (
    EntriesSection,
    Entry,
    ListSection,
    Pair,
    PairsSection,
    Profile,
    ProfileHeader,
    ProfileLink,
    TagCategory,
    TagsSection,
    TextSection,
)
from services.docx_export import build_resume_document


def _square_picture_data_url(size: int = 400) -> str:
    """A square test photo — deliberately the wrong aspect ratio for the
    picture frame (~1.11:1), to exercise the center-crop path."""
    image = Image.new("RGB", (size, size), color=(200, 100, 50))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    encoded = base64.b64encode(buffer.getvalue()).decode()
    return f"data:image/png;base64,{encoded}"


def _make_profile(**overrides) -> Profile:
    defaults = dict(
        id="1",
        name="Test Profile",
        header=ProfileHeader(
            full_name="Jane Doe",
            career_title="AI Engineer",
            email="jane@example.com",
            phone="555-0100",
            location="Remote",
            links=[ProfileLink(id="l1", label="GitHub", url="https://github.com")],
            primary_color="#2563eb",
            secondary_color="#7c3aed",
        ),
        sections=[],
    )
    defaults.update(overrides)
    return Profile(**defaults)


def _reload(document: Document) -> Document:
    """Round-trips through bytes, the way the real export endpoint would
    serve it, to make sure the generated XML is actually valid."""
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return Document(buffer)


def _paragraph_texts(document: Document) -> list[str]:
    return [p.text for p in document.paragraphs]


def _header_cell(document: Document):
    """The header banner (name/title/contact/links) lives in a table cell
    cloned from the template, not in the document's top-level paragraphs —
    the last cell either way, since the no-picture layout is a single
    full-width cell while the picture layout has a picture column first."""
    return document.tables[0].rows[0].cells[-1]


def test_includes_header_fields() -> None:
    profile = _make_profile()
    document = _reload(build_resume_document(profile))

    header_texts = [p.text for p in _header_cell(document).paragraphs]
    assert "Jane Doe" in header_texts
    assert "AI Engineer" in header_texts
    assert any("jane@example.com" in t for t in header_texts)


def test_contact_line_breaks_before_location_instead_of_joining_with_a_pipe() -> None:
    profile = _make_profile()
    document = build_resume_document(profile)

    contact_paragraph = _header_cell(document).paragraphs[2]
    runs = contact_paragraph._p.findall(qn("w:r"))

    has_break = any(run.find(qn("w:br")) is not None for run in runs)
    assert has_break

    run_texts = [
        run.find(qn("w:t")).text for run in runs if run.find(qn("w:t")) is not None
    ]
    assert run_texts == ["jane@example.com | 555-0100", "Remote"]


def test_header_name_is_white_against_the_colored_banner() -> None:
    profile = _make_profile()
    document = build_resume_document(profile)

    name_paragraph = _header_cell(document).paragraphs[0]
    assert str(name_paragraph.runs[0].font.color.rgb) == "FFFFFF"


def test_header_banner_is_filled_with_the_primary_color() -> None:
    profile = _make_profile()
    document = build_resume_document(profile)

    table = document.tables[0]
    for cell in table.rows[0].cells:
        shd = cell._tc.tcPr.find(qn("w:shd"))
        assert shd.get(qn("w:fill")) == "2563EB"


def test_header_link_is_a_real_hyperlink() -> None:
    profile = _make_profile()
    document = _reload(build_resume_document(profile))

    links_paragraph = next(
        p for p in _header_cell(document).paragraphs if "GitHub" in p.text
    )
    hyperlinks = links_paragraph._p.findall(qn("w:hyperlink"))
    assert len(hyperlinks) == 1


def test_header_link_with_an_unsafe_scheme_is_not_a_hyperlink() -> None:
    profile = _make_profile(
        header=ProfileHeader(
            full_name="Jane Doe",
            career_title="AI Engineer",
            email="jane@example.com",
            phone="555-0100",
            location="Remote",
            links=[ProfileLink(id="l1", label="Evil", url="javascript:alert(1)")],
            primary_color="#2563eb",
            secondary_color="#7c3aed",
        )
    )
    document = _reload(build_resume_document(profile))

    links_paragraph = next(
        p for p in _header_cell(document).paragraphs if "Evil" in p.text
    )
    hyperlinks = links_paragraph._p.findall(qn("w:hyperlink"))
    assert len(hyperlinks) == 0


def test_header_with_no_links_renders_an_empty_links_paragraph() -> None:
    profile = _make_profile(
        header=ProfileHeader(
            full_name="Jane Doe",
            career_title="AI Engineer",
            email="jane@example.com",
            phone="555-0100",
            location="Remote",
            links=[],
            primary_color="#2563eb",
            secondary_color="#7c3aed",
        )
    )
    document = _reload(build_resume_document(profile))

    links_paragraph = _header_cell(document).paragraphs[3]
    assert links_paragraph.text == ""


def test_unused_template_hyperlink_relationships_are_removed() -> None:
    profile = _make_profile()
    document = _reload(build_resume_document(profile))

    targets = {
        rel.target_ref
        for rel in document.part.rels.values()
        if "hyperlink" in rel.reltype
    }
    # Only the Profile's own link should remain — not the template's
    # placeholder example links (e.g. github.com/yourusername).
    assert targets == {"https://github.com"}


def test_text_section_renders_bold_run_from_markdown_lite() -> None:
    profile = _make_profile(
        sections=[
            TextSection(id="s1", type="text", title="Summary", body="**Bold** text"),
        ]
    )
    document = _reload(build_resume_document(profile))

    body_paragraph = next(p for p in document.paragraphs if "Bold" in p.text)
    bold_run = next(r for r in body_paragraph.runs if r.text == "Bold")
    assert bold_run.bold is True


def test_text_section_link_becomes_a_real_hyperlink() -> None:
    profile = _make_profile(
        sections=[
            TextSection(
                id="s1",
                type="text",
                title="Summary",
                body="See [my site](https://example.com)",
            ),
        ]
    )
    document = _reload(build_resume_document(profile))

    body_paragraph = next(p for p in document.paragraphs if "my site" in p.text)
    hyperlinks = body_paragraph._p.findall(qn("w:hyperlink"))
    assert len(hyperlinks) == 1


def test_entries_section_renders_heading_dates_on_separate_lines_and_body() -> None:
    profile = _make_profile(
        sections=[
            EntriesSection(
                id="s1",
                type="entries",
                title="Experience",
                entries=[
                    Entry(
                        id="e1",
                        heading="Engineer, Acme",
                        dates="2020 - 2022",
                        body="Did the work.",
                    ),
                ],
            ),
        ]
    )
    document = _reload(build_resume_document(profile))
    texts = _paragraph_texts(document)

    # Matches the template's own design: the dates sit on their own line,
    # not crammed inline after the job title.
    assert "Engineer, Acme" in texts
    assert "2020 - 2022" in texts
    assert texts.index("2020 - 2022") == texts.index("Engineer, Acme") + 1
    assert any(t == "Did the work." for t in texts)


def test_entries_section_job_title_and_dates_use_template_typography() -> None:
    profile = _make_profile(
        sections=[
            EntriesSection(
                id="s1",
                type="entries",
                title="Experience",
                entries=[
                    Entry(
                        id="e1",
                        heading="Engineer, Acme",
                        dates="2020 - 2022",
                        body="",
                    ),
                ],
            ),
        ]
    )
    document = build_resume_document(profile)

    title_paragraph = next(p for p in document.paragraphs if p.text == "Engineer, Acme")
    dates_paragraph = next(p for p in document.paragraphs if p.text == "2020 - 2022")

    assert title_paragraph.runs[0].bold is True
    assert str(dates_paragraph.runs[0].font.color.rgb) == "555555"


def test_entries_section_omits_dates_line_when_dates_is_empty() -> None:
    profile = _make_profile(
        sections=[
            EntriesSection(
                id="s1",
                type="entries",
                title="Education",
                entries=[
                    Entry(id="e1", heading="Degree, University", dates="", body=""),
                ],
            ),
        ]
    )
    document = _reload(build_resume_document(profile))
    texts = _paragraph_texts(document)

    index = texts.index("Degree, University")
    # No stray empty dates/body paragraph — either nothing follows the job
    # title at all, or whatever comes next is real content, not "".
    assert index == len(texts) - 1 or texts[index + 1] != ""


def test_list_section_renders_each_item_as_its_own_paragraph() -> None:
    profile = _make_profile(
        sections=[
            ListSection(id="s1", type="list", title="Projects", items=["A", "B"]),
        ]
    )
    document = _reload(build_resume_document(profile))
    texts = _paragraph_texts(document)

    assert "A" in texts
    assert "B" in texts


def test_tags_section_renders_bold_label_and_plain_joined_tags() -> None:
    profile = _make_profile(
        sections=[
            TagsSection(
                id="s1",
                type="tags",
                title="Skills",
                categories=[
                    TagCategory(id="c1", label="Languages", items=["Python", "TS"]),
                ],
            ),
        ]
    )
    document = _reload(build_resume_document(profile))

    tags_paragraph = next(p for p in document.paragraphs if "Languages" in p.text)
    assert tags_paragraph.text == "Languages: Python, TS"
    assert tags_paragraph.runs[0].text == "Languages:"
    assert tags_paragraph.runs[0].bold is True
    assert tags_paragraph.runs[1].bold is not True


def test_pairs_section_renders_label_colon_value() -> None:
    profile = _make_profile(
        sections=[
            PairsSection(
                id="s1",
                type="pairs",
                title="Languages",
                pairs=[Pair(id="p1", left="English", right="Fluent")],
            ),
        ]
    )
    document = _reload(build_resume_document(profile))
    texts = _paragraph_texts(document)

    assert "English: Fluent" in texts


def test_section_title_uses_the_secondary_color() -> None:
    profile = _make_profile(
        sections=[
            TextSection(id="s1", type="text", title="Summary", body="hi"),
        ]
    )
    document = build_resume_document(profile)

    heading_paragraph = next(p for p in document.paragraphs if p.text == "Summary")
    assert str(heading_paragraph.runs[0].font.color.rgb) == "7C3AED"


def test_page_margins_and_size_match_the_template() -> None:
    profile = _make_profile()
    document = build_resume_document(profile)

    section = document.sections[0]
    assert section.left_margin == section.right_margin == 457200
    assert section.top_margin == section.bottom_margin == 457200


def test_no_picture_uses_the_single_column_header_layout() -> None:
    profile = _make_profile()
    document = build_resume_document(profile)

    assert len(document.tables[0].columns) == 1


def test_picture_uses_the_two_column_header_layout() -> None:
    profile = _make_profile(
        header=ProfileHeader(
            full_name="Jane Doe",
            career_title="AI Engineer",
            email="jane@example.com",
            phone="555-0100",
            location="Remote",
            links=[],
            primary_color="#2563eb",
            secondary_color="#7c3aed",
            picture=_square_picture_data_url(),
        )
    )
    document = _reload(build_resume_document(profile))

    assert len(document.tables[0].columns) == 2


def test_picture_is_cropped_to_the_frames_exact_aspect_ratio() -> None:
    profile = _make_profile(
        header=ProfileHeader(
            full_name="Jane Doe",
            career_title="AI Engineer",
            email="jane@example.com",
            phone="555-0100",
            location="Remote",
            links=[],
            primary_color="#2563eb",
            secondary_color="#7c3aed",
            picture=_square_picture_data_url(),
        )
    )
    document = _reload(build_resume_document(profile))

    drawing = next(document.element.body.iter(qn("w:drawing")))
    extent = next(drawing.iter(qn("wp:extent")))
    target_ratio = int(extent.get("cx")) / int(extent.get("cy"))

    blip = next(drawing.iter(qn("a:blip")))
    r_id = blip.get(qn("r:embed"))
    embedded_image = document.part.rels[r_id].target_part.image
    embedded_ratio = embedded_image.px_width / embedded_image.px_height

    assert embedded_ratio == pytest.approx(target_ratio, rel=0.01)
    # The frame's own size is untouched regardless of what was uploaded —
    # that's what "exact same width and height" in the frame means.
    assert extent.get("cx") == "1504543"
    assert extent.get("cy") == "1357200"


def test_picture_crop_replaces_the_templates_own_srcrect() -> None:
    profile = _make_profile(
        header=ProfileHeader(
            full_name="Jane Doe",
            career_title="AI Engineer",
            email="jane@example.com",
            phone="555-0100",
            location="Remote",
            links=[],
            primary_color="#2563eb",
            secondary_color="#7c3aed",
            picture=_square_picture_data_url(),
        )
    )
    document = _reload(build_resume_document(profile))

    drawing = next(document.element.body.iter(qn("w:drawing")))
    blip_fill = next(drawing.iter(qn("pic:blipFill")))
    # The template's own srcRect crop is tuned to its placeholder photo and
    # would distort a different one — the uploaded photo is already cropped
    # to the right aspect ratio in Python, so no XML-level crop should apply.
    assert blip_fill.find(qn("a:srcRect")) is None


def test_unused_placeholder_image_relationship_is_removed_without_a_picture() -> None:
    profile = _make_profile()
    document = _reload(build_resume_document(profile))

    image_targets = [
        rel.target_ref for rel in document.part.rels.values() if "image" in rel.reltype
    ]
    assert image_targets == []


def test_section_heading_has_a_bottom_border_matching_the_secondary_color() -> None:
    profile = _make_profile(
        sections=[
            TextSection(id="s1", type="text", title="Summary", body="hi"),
        ]
    )
    document = _reload(build_resume_document(profile))

    heading = next(p for p in document.paragraphs if p.text == "Summary")
    p_bdr = heading._p.pPr.find(qn("w:pBdr"))
    bottom = p_bdr.find(qn("w:bottom"))

    assert bottom.get(qn("w:val")) == "single"
    assert bottom.get(qn("w:color")) == "7C3AED"


def test_cloned_paragraphs_do_not_carry_the_templates_leftover_div_id() -> None:
    """The template's own paragraphs each reference a `w:divId`, which
    links to a *separate* border definition in word/webSettings.xml (not
    the paragraph's own w:pBdr). Left in place, this rendered as a second,
    uncontrollable-colored line under every heading and body paragraph in
    Word, invisible to any check that only looks at the paragraph itself.
    """
    profile = _make_profile(
        sections=[
            TextSection(id="s1", type="text", title="Summary", body="Body text."),
            EntriesSection(
                id="s2",
                type="entries",
                title="Experience",
                entries=[
                    Entry(
                        id="e1",
                        heading="Engineer, Acme",
                        dates="2020 - 2022",
                        body="- A bullet.",
                    ),
                ],
            ),
            TagsSection(
                id="s3",
                type="tags",
                title="Skills",
                categories=[TagCategory(id="c1", label="Languages", items=["Python"])],
            ),
        ]
    )
    document = _reload(build_resume_document(profile))

    for paragraph in document.paragraphs:
        pPr = paragraph._p.pPr
        div_id = pPr.find(qn("w:divId")) if pPr is not None else None
        assert div_id is None, f"stray divId on {paragraph.text!r}"

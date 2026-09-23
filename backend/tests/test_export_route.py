import io

from docx import Document
from fastapi.testclient import TestClient

import main
from auth.auth import get_current_user
from auth.subscription import require_subscription_plan
from db.models import SubscriptionPlan, User

fake_user = User(id=1, clerk_id="test", email="test@example.com", role="user")
main.app.dependency_overrides[get_current_user] = lambda: fake_user
main.app.dependency_overrides[require_subscription_plan(SubscriptionPlan.BASIC)] = (
    lambda: None
)

client = TestClient(main.app)


def _profile_payload() -> dict:
    return {
        "id": "1",
        "name": "Test Profile",
        "header": {
            "fullName": "Jane Doe",
            "careerTitle": "AI Engineer",
            "email": "jane@example.com",
            "phone": "555-0100",
            "location": "Remote",
            "links": [],
            "primaryColor": "#2563eb",
            "secondaryColor": "#7c3aed",
        },
        "sections": [
            {
                "id": "s1",
                "type": "text",
                "title": "Summary",
                "body": "**Bold** summary.",
            }
        ],
    }


def test_export_returns_a_valid_docx_with_the_profile_content() -> None:
    response = client.post("/profile/export", json=_profile_payload())

    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    assert "Jane_Doe.docx" in response.headers["content-disposition"]

    document = Document(io.BytesIO(response.content))
    header_texts = [p.text for p in document.tables[0].rows[0].cells[1].paragraphs]
    assert "Jane Doe" in header_texts
    assert any(p.text == "Summary" for p in document.paragraphs)
